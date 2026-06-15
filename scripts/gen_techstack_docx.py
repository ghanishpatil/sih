"""Generate a polished TECH_STACK.docx for the SKH Platform."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ─── Brand palette ───────────────────────────────────────────────
BRAND = RGBColor(0x4F, 0x46, 0xE5)      # indigo
BRAND_DARK = RGBColor(0x31, 0x2E, 0x81)
CYAN = RGBColor(0x06, 0xB6, 0xD4)
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEADER_BG = "4F46E5"
ROW_ALT = "EEF2FF"
SEC_BG = "ECFDF5"

doc = Document()

# Base style
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.font.color.rgb = INK


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_text(cell, text, *, bold=False, color=None, size=10.5, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    return run


def add_heading_bar(text, subtitle=None):
    """A colored section heading."""
    p = doc.add_paragraph()
    p.space_before = Pt(14)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = BRAND
    # bottom border
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "4F46E5")
    pbdr.append(bottom)
    pPr.append(pbdr)
    if subtitle:
        s = doc.add_paragraph()
        sr = s.add_run(subtitle)
        sr.italic = True
        sr.font.size = Pt(9.5)
        sr.font.color.rgb = MUTED


def add_two_col_table(rows, headers=("Component", "Technology")):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_text(hdr[i], h, bold=True, color=WHITE, size=10.5)
        shade_cell(hdr[i], HEADER_BG)
    for idx, (a, b) in enumerate(rows):
        cells = table.add_row().cells
        set_cell_text(cells[0], a, bold=True, color=INK)
        set_cell_text(cells[1], b, color=INK)
        if idx % 2 == 1:
            shade_cell(cells[0], ROW_ALT)
            shade_cell(cells[1], ROW_ALT)
    return table


def add_three_col_table(rows, headers):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_text(hdr[i], h, bold=True, color=WHITE, size=10)
        shade_cell(hdr[i], HEADER_BG)
    for idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], val, bold=(i == 0), color=INK, size=9.5)
            if idx % 2 == 1:
                shade_cell(cells[i], ROW_ALT)
    return table


def add_security_block(title, items):
    p = doc.add_paragraph()
    p.space_before = Pt(8)
    run = p.add_run("  " + title)
    run.bold = True
    run.font.size = Pt(11.5)
    run.font.color.rgb = BRAND_DARK
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), SEC_BG)
    pPr.append(shd)
    for it in items:
        b = doc.add_paragraph(style="List Bullet")
        br = b.add_run(it)
        br.font.size = Pt(10)
        br.font.color.rgb = INK


# ─── Title page ──────────────────────────────────────────────────
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
tr = title.add_run("Smart Kopargaon Hackathon")
tr.bold = True
tr.font.size = Pt(28)
tr.font.color.rgb = BRAND

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sr = sub.add_run("Platform Technology Stack & Security Overview")
sr.font.size = Pt(14)
sr.font.color.rgb = MUTED

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
mr = meta.add_run("Full-stack hackathon management platform  •  Sanjivani University")
mr.italic = True
mr.font.size = Pt(10)
mr.font.color.rgb = MUTED

doc.add_paragraph()

# ─── Frontend ────────────────────────────────────────────────────
add_heading_bar("Frontend")
add_two_col_table([
    ("Framework", "React 19"),
    ("Build tool", "Vite 8"),
    ("Routing", "React Router DOM v7 (SPA, lazy-loaded routes)"),
    ("Styling", "Tailwind CSS 3"),
    ("UI primitives", "Radix UI + shadcn/ui"),
    ("Animations", "Framer Motion"),
    ("Smooth scroll", "Lenis"),
    ("Icons", "Lucide React"),
    ("Charts / analytics", "Recharts"),
    ("Tables", "TanStack React Table"),
    ("Forms", "React Hook Form"),
    ("Validation", "Zod"),
    ("State management", "Zustand + React Context"),
    ("Rich animations", "Rive + Lottie"),
    ("Firebase client", "Firebase JS SDK v11 (Auth, Firestore, Storage)"),
])

# ─── Backend ─────────────────────────────────────────────────────
add_heading_bar("Backend")
add_two_col_table([
    ("Runtime", "Node.js (ES Modules)"),
    ("Web framework", "Express 5"),
    ("Database", "Cloud Firestore (Firebase Admin SDK v13)"),
    ("Authentication", "Firebase Authentication (JWT + custom claims)"),
    ("File storage", "Firebase Storage"),
    ("Payments", "Razorpay (REST API + webhooks)"),
    ("Email", "Brevo (Sendinblue) REST API"),
    ("AI chatbot", "Google Gemini (gemini-2.5-flash)"),
    ("File uploads", "Multer (memory storage)"),
    ("Logging", "Morgan"),
    ("Compression", "compression (gzip/brotli)"),
])

# ─── Infrastructure ──────────────────────────────────────────────
add_heading_bar("Infrastructure & Deployment")
add_two_col_table([
    ("Firebase project", "verify-sih"),
    ("Frontend hosting", "Vercel"),
    ("Backend hosting", "Railway"),
    ("Source control", "GitHub"),
    ("Database rules", "Firestore Security Rules"),
    ("Storage rules", "Firebase Storage Rules"),
])

doc.add_page_break()

# ─── Feature mapping ─────────────────────────────────────────────
add_heading_bar("Feature → Technology Mapping")
add_three_col_table([
    ("Auth & roles", "Firebase Auth + React Context", "Admin SDK, custom claims"),
    ("Team formation", "React Hook Form + Zod", "Firestore transactions, crypto invite codes"),
    ("Registration & payment", "React + Razorpay checkout", "Razorpay orders/webhooks, transactions"),
    ("Submissions", "React forms", "Storage URL validation, Firestore"),
    ("Real-time chat", "Firestore onSnapshot", "Express + Admin SDK writes"),
    ("Multi-phase competition", "EventContext (onSnapshot)", "Phase state machine in Firestore"),
    ("Jury evaluation", "React + Recharts", "Normalized scoring, Firestore"),
    ("Skill tags & profiles", "SkillTagEditor (Framer Motion)", "Admin SDK writes to users"),
    ("Progress tracker", "Framer Motion stepper", "Derived client-side"),
    ("Team matchmaking", "Tabs + request flow", "joinRequests collection, transactions"),
    ("Admin AI chatbot", "React chat UI", "Google Gemini + service account"),
    ("Security center", "React dashboards", "securityEvents, incidents, honeypots"),
    ("Server caching", "—", "In-memory response cache"),
], headers=("Feature", "Frontend", "Backend / Data"))

doc.add_page_break()

# ─── Security ────────────────────────────────────────────────────
add_heading_bar("Security Features by Layer")

add_security_block("Authentication & Authorization", [
    "Firebase Authentication with JWT bearer tokens, verified server-side on every request",
    "Role-based access control (admin, participant, judge, mentor, banned) via requireRole middleware",
    "Roles stored as tamper-proof custom claims in the JWT, with Firestore fallback",
    "Banned status always re-checked from Firestore to prevent stale JWTs granting access",
    "Self-service profile updates restricted to a field allowlist",
])

add_security_block("Firestore Security Rules", [
    "Default deny-all catch-all rule",
    "Most collections are server-only writes (Admin SDK bypasses rules; clients blocked)",
    "Per-document ownership checks (users read only their own data; judges see only assigned teams)",
    "Event documents restricted to admin reads — public data served via curated API snapshots",
    "joinRequests, securityEvents, securityIncidents, webhookLog, activityLog: admin-read, server-only write",
])

add_security_block("Payment Security (Razorpay)", [
    "HMAC-SHA256 signature verification on both client callbacks and webhooks",
    "Timing-safe signature comparison (constant-time, fixed-length buffers)",
    "Idempotency locks (webhookEvents/{paymentId}) to dedupe webhook retries",
    "Payment writes wrapped in Firestore transactions to prevent double-processing",
    "Order amount/currency validated against event config before marking paid",
])

add_security_block("Input Validation & Sanitization", [
    "Firestore document ID validation (no slashes, null bytes, or path traversal)",
    "Submission URLs restricted to Firebase Storage hosts scoped to the team's own path",
    "Input length caps on all user-provided strings (invite codes, skills, bio, messages)",
    "HTML entity escaping in email templates to prevent XSS",
    "File upload type/size restrictions via Multer filters",
])

add_security_block("Network & Transport Security", [
    "Helmet security headers (with cross-origin resource policy)",
    "Strict CORS allowlist (production rejects unknown origins and no-Origin requests to prevent CSRF)",
    "Rate limiting — global API limiter plus tighter per-endpoint limits for payments, team actions, webhooks",
    "HTTPS enforced for all external URLs",
])

add_security_block("Threat Detection & Monitoring", [
    "55+ honeypot trap routes that auto-log security events and create incidents",
    "Append-only security event logging with IP, device fingerprint, and GeoIP enrichment",
    "Incident management system with deduplication",
    "Platform-wide activity feed and audit logs for all admin actions",
    "User-Agent / bot scanner detection",
])

add_security_block("Privacy", [
    "Emails never exposed in matchmaking or to non-leader teammates",
    "PII referenced by key name, not value, in logs",
    "Secrets (API keys, service accounts) kept in environment variables, never committed",
    "Error responses suppress stack traces in production",
])

add_security_block("Operational Safety", [
    "Server-side caching with TTL, max-size caps, and periodic eviction (prevents memory leaks)",
    "Environment validation on startup (warns on misconfigured CORS, email, webhook secrets)",
    "Fire-and-forget logging never crashes the main request path",
    "Graceful degradation when external services (email, GeoIP, Razorpay) are unavailable",
])

# Footer note
doc.add_paragraph()
foot = doc.add_paragraph()
foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = foot.add_run("Last updated: June 15, 2026")
fr.italic = True
fr.font.size = Pt(9)
fr.font.color.rgb = MUTED

import os
out = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "TECH_STACK.docx")
doc.save(out)
print("Saved:", out)
